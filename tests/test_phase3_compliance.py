"""第 3 期合规:高风险功能开关、公开链接过期、授权留痕、AI 生成内容标识.

全部走真实函数 + 临时 SQLite,不依赖 fastapi(接口层只是把 FeatureDisabled
转成 403,这里直接验证下层拒绝与返回的大白话原因)。
"""
import asyncio
import io
import json
import os
import shutil
import subprocess
import tempfile
import time
import unittest
from unittest import mock

from app import auth, avatar, db, features, growth, imagehunt, linkgrab, matrixpub, mplayout


class _DbCase(unittest.TestCase):
    _template_dir = None

    @classmethod
    def _template(cls) -> str:
        """迁移一次建好的空库，每个用例复制一份(省去每次全量建表)。"""
        if _DbCase._template_dir is None:
            _DbCase._template_dir = tempfile.mkdtemp(prefix="phase3-template-")
            old = db.DB_PATH
            if db._conn is not None:
                db._conn.close()
            db._conn = None
            db.DB_PATH = os.path.join(_DbCase._template_dir, "template.db")
            db.conn()
            db.conn().execute("PRAGMA wal_checkpoint(TRUNCATE)")
            db._close_all_connections()
            db._conn = None
            db._conn_path = None
            db.DB_PATH = old
        return os.path.join(_DbCase._template_dir, "template.db")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old_db_path = db.DB_PATH
        self.old_public_dir = avatar.PUBLIC_DIR
        template = self._template()
        if db._conn is not None:
            db._conn.close()
        db._conn = None
        db.DB_PATH = os.path.join(self.tmp.name, "phase3.db")
        shutil.copyfile(template, db.DB_PATH)
        avatar.PUBLIC_DIR = os.path.join(self.tmp.name, "public")
        os.makedirs(avatar.PUBLIC_DIR, exist_ok=True)
        db.conn()
        db.insert("tenants", {"id": 1, "name": "平台", "balance": 0})
        db.insert("tenants", {"id": 2, "name": "火锅店", "balance": 30})
        db.insert("tenants", {"id": 3, "name": "奶茶店", "balance": 30})
        self.secret = mock.patch.object(auth, "_secret", return_value=b"k" * 64)
        self.secret.start()
        self.as_tenant(2)

    def tearDown(self):
        self.secret.stop()
        auth.set_current(None)
        if db._conn is not None:
            db._conn.close()
        db._conn = None
        db.DB_PATH = self.old_db_path
        avatar.PUBLIC_DIR = self.old_public_dir
        self.tmp.cleanup()

    def as_tenant(self, tid, role="owner"):
        auth.set_current({"id": tid * 10, "tenant_id": tid, "username": f"老板{tid}",
                          "role": role, "modules": ["avatar", "content"]})


# ---------------------------------------------------------------- 1. 开关
class FeatureSwitchTests(_DbCase):
    def test_all_high_risk_features_default_off_with_risk_text(self):
        for key in ("matrix_autopub", "imagehunt", "linkgrab_video", "lead_search_scrape"):
            self.assertFalse(features.is_enabled(key), key)
            with self.assertRaises(features.FeatureDisabled) as caught:
                features.require(key)
            self.assertEqual(403, caught.exception.status_code)
            self.assertTrue(str(caught.exception))
        overview = features.admin_overview()
        self.assertEqual(4, len(overview["features"]))
        for item in overview["features"]:
            self.assertFalse(item["default"])
            self.assertFalse(item["enabled"])
            self.assertGreater(len(item["risk"]), 20)

    def test_platform_switch_and_root_only_tenant_override(self):
        features.set_platform("imagehunt", True)
        self.assertTrue(features.is_enabled("imagehunt"))
        # 企业级覆盖:只关掉 2 号企业,3 号仍跟随平台
        features.set_tenant("imagehunt", 2, False)
        self.assertFalse(features.is_enabled("imagehunt"))
        self.assertTrue(features.is_enabled("imagehunt", 3))
        features.set_tenant("imagehunt", 2, None)
        self.assertTrue(features.is_enabled("imagehunt"))
        # 平台关、单独给 3 号开
        features.set_platform("imagehunt", False)
        features.set_tenant("imagehunt", 3, True)
        self.assertTrue(features.is_enabled("imagehunt", 3))
        self.assertFalse(features.is_enabled("imagehunt", 2))
        overrides = features.admin_overview()["features"][1]["overrides"]
        self.assertEqual([{"tenant_id": 3, "enabled": True}], overrides)
        # 后台任务没有登录上下文时只看平台开关
        auth.set_current(None)
        self.assertFalse(features.is_enabled("imagehunt"))

    def test_matrix_autopub_backend_refuses_bind_and_enqueue(self):
        with self.assertRaises(features.FeatureDisabled) as caught:
            matrixpub.add_account(2, "xhs", "主号", "sessionid=" + "c" * 60)
        self.assertIn("半自动发布", str(caught.exception))
        with self.assertRaises(features.FeatureDisabled):
            matrixpub.enqueue(2, "xhs", "a1", {"title": "t"})
        with self.assertRaises(features.FeatureDisabled):
            asyncio.run(matrixpub.check_account(2, "a1"))
        self.assertEqual(0, db.one("SELECT COUNT(*) n FROM pub_task")["n"])

    def test_imagehunt_refuses_search_download_and_pipeline(self):
        with mock.patch.object(imagehunt, "_bing") as bing:
            with self.assertRaises(features.FeatureDisabled) as caught:
                asyncio.run(imagehunt.search("火锅 实拍"))
            with self.assertRaises(features.FeatureDisabled):
                asyncio.run(imagehunt.fetch_image("https://example.com/a.jpg"))
        bing.assert_not_called()
        self.assertIn("AI", str(caught.exception))
        job_id = db.insert("job", {"tenant_id": 2, "brief_json": "{}"})
        steps = []
        with mock.patch.object(imagehunt, "search") as search:
            images = asyncio.run(imagehunt.hunt_for_job(
                job_id, 1, "火锅", [{"slot": "封面"}], 1,
                progress=lambda kind, label="": steps.append((kind, label))))
        self.assertEqual([], images)
        search.assert_not_called()
        self.assertIn("全网抓图", steps[0][1])

    def test_imagehunt_tenant_override_allows_search(self):
        features.set_tenant("imagehunt", 2, True)

        async def fake(_cli, _q, _n):
            return [{"img": "https://img.example/1.jpg", "thumb": "", "page": "", "from": "bing"}]

        with mock.patch.object(imagehunt, "_bing", new=fake), \
                mock.patch.object(imagehunt, "_baidu", new=fake), \
                mock.patch.object(imagehunt, "_so360", new=fake):
            items = asyncio.run(imagehunt.search("火锅", 4))
        self.assertEqual(1, len(items))

    def test_video_link_transcription_refused_but_articles_allowed(self):
        with self.assertRaises(features.FeatureDisabled) as caught:
            linkgrab.ensure_video_allowed("https://v.douyin.com/abc/")
        self.assertIn("请上传你自己的视频/音频文件", str(caught.exception))
        # 公众号/普通网页文章不受这个开关影响
        linkgrab.ensure_video_allowed("https://mp.weixin.qq.com/s/abc")
        features.set_platform("linkgrab_video", True)
        linkgrab.ensure_video_allowed("https://v.douyin.com/abc/")

    def test_lead_radar_skips_search_engine_scraping_and_explains(self):
        steps = []
        with mock.patch.object(growth, "_public_lead_search") as scrape:
            rows = asyncio.run(growth.direct_lead_sources(
                "餐饮", "成都", "火锅", tenant_id=2,
                progress=lambda kind, label="": steps.append(label)))
        self.assertEqual([], rows)
        scrape.assert_not_called()
        self.assertIn("搜索引擎", steps[0])

    def test_queued_auto_publish_is_closed_with_semi_auto_way_out(self):
        pid = db.insert("pub_task", {
            "tenant_id": 2, "platform": "xhs", "account": "a1",
            "payload_json": json.dumps({"title": "新品"}), "status": "queued",
            "submission_state": "not_submitted",
        })
        matrixpub._PUB_SEM = None
        try:
            with mock.patch.object(matrixpub, "_run_task_inner") as inner:
                asyncio.run(matrixpub.run_task(pid))
        finally:
            matrixpub._PUB_SEM = None
        inner.assert_not_called()
        row = db.one("SELECT status,fail_json,submission_state FROM pub_task WHERE id=?", (pid,))
        self.assertEqual("failed", row["status"])
        self.assertEqual("not_submitted", row["submission_state"])
        fail = json.loads(row["fail_json"])
        self.assertEqual("disabled", fail["kind"])
        self.assertIn("半自动发布", fail["fix"])
        self.assertEqual("https://creator.xiaohongshu.com", fail["home"])

    def test_existing_accounts_still_listable_and_deletable_when_off(self):
        features.set_platform("matrix_autopub", True)
        acc = matrixpub.add_account(2, "xhs", "主号", "sessionid=" + "c" * 60)
        features.set_platform("matrix_autopub", False)
        self.assertEqual([acc["id"]], [a["id"] for a in matrixpub.pub_list(2)])
        matrixpub.del_account(2, acc["id"])
        self.assertEqual([], matrixpub.pub_list(2))

    def test_compliance_config_is_bounded(self):
        self.assertEqual(7, features.config_int("pubfile_ttl_days"))
        features.set_config({"pubfile_ttl_days": "3"})
        self.assertEqual(3, features.config_int("pubfile_ttl_days"))
        with self.assertRaises(ValueError):
            features.set_config({"pubfile_ttl_days": "999"})
        features.set_config({"pubfile_ttl_days": ""})
        self.assertEqual(7, features.config_int("pubfile_ttl_days"))


# ---------------------------------------------------------------- 2. 公开链接
class PubfileSignatureTests(_DbCase):
    REL = "job7/media_v1_0.png"

    def _token(self, url):
        return url.split("/")[2]

    def test_signed_link_expires_after_configured_days(self):
        now = 1_800_000_000
        url = mplayout.sign_file(self.REL, now=now)
        token = self._token(url)
        self.assertRegex(token, r"^\d+-[0-9a-f]{64}$")
        self.assertTrue(url.endswith("/" + self.REL))
        self.assertTrue(mplayout.verify_file(token, self.REL, now=now + 6 * 86400))
        self.assertTrue(mplayout.verify_file(token, self.REL, now=now + 7 * 86400))
        self.assertFalse(mplayout.verify_file(token, self.REL, now=now + 9 * 86400))
        # 同一天重复生成完全相同(草稿去重哈希不抖动)
        self.assertEqual(url, mplayout.sign_file(self.REL, now=now + 60))
        features.set_config({"pubfile_ttl_days": 1})
        short = self._token(mplayout.sign_file(self.REL, now=now))
        self.assertFalse(mplayout.verify_file(short, self.REL, now=now + 3 * 86400))

    def test_tampered_signature_path_or_expiry_rejected(self):
        now = 1_800_000_000
        token = self._token(mplayout.sign_file(self.REL, now=now))
        expires, sig = token.split("-")
        flipped = ("0" if sig[0] != "0" else "1") + sig[1:]
        self.assertFalse(mplayout.verify_file(f"{expires}-{flipped}", self.REL, now=now))
        self.assertFalse(mplayout.verify_file(token, "job8/media_v1_0.png", now=now))
        self.assertFalse(mplayout.verify_file(f"{int(expires) + 86400}-{sig}", self.REL, now=now))
        self.assertFalse(mplayout.verify_file(sig[:20], self.REL, now=now))
        self.assertFalse(mplayout.verify_file("", self.REL, now=now))

    def test_legacy_permanent_signature_only_valid_during_transition(self):
        start = 1_800_000_000
        features.legacy_since(now=start)
        legacy = mplayout._legacy_sig(self.REL)
        self.assertTrue(mplayout.verify_file(legacy, self.REL, now=start + 6 * 86400))
        self.assertFalse(mplayout.verify_file(legacy, self.REL, now=start + 8 * 86400))
        tampered = legacy[:-1] + ("1" if legacy[-1] == "0" else "0")
        self.assertFalse(mplayout.verify_file(tampered, self.REL, now=start))
        # 过渡期起点一经记录不再后移
        self.assertEqual(start, features.legacy_since(now=start + 100 * 86400))
        features.set_config({"pubfile_legacy_days": 0})
        self.assertFalse(mplayout.verify_file(legacy, self.REL, now=start + 60))


class PublicAvatarLinkTests(_DbCase):
    def _file(self, name, age_days=0.0, data=b"x"):
        path = os.path.join(avatar.PUBLIC_DIR, name)
        with open(path, "wb") as handle:
            handle.write(data)
        stamp = time.time() - age_days * 86400
        os.utime(path, (stamp, stamp))
        return path

    def test_vendor_link_is_signed_temporary_and_tamper_proof(self):
        name = "a" * 32 + ".jpg"
        path = self._file(name)
        now = 1_800_000_000
        url = avatar.signed_public_url(name, now=now)
        self.assertTrue(url.startswith("https://paihuo.ai/pub/s/"))
        _, expires, sig, got = url.rsplit("/", 3)
        self.assertEqual(name, got)
        self.assertEqual(now + 6 * 3600, int(expires))
        self.assertEqual(os.path.realpath(path),
                         avatar.resolve_signed_public(expires, sig, name, now=now + 60))
        self.assertIsNone(avatar.resolve_signed_public(expires, sig, name, now=now + 7 * 3600))
        self.assertIsNone(avatar.resolve_signed_public(int(expires) + 1, sig, name, now=now))
        self.assertIsNone(avatar.resolve_signed_public(expires, sig, "b" * 32 + ".jpg", now=now))
        self.assertIsNone(avatar.resolve_signed_public(expires, "0" * 64, name, now=now))
        self.assertIsNone(avatar.resolve_signed_public(expires, sig, "../" + name, now=now))
        # UUID 素材不再能按文件名直接取;只有固定前缀的宣传素材可以
        self.assertIsNone(avatar.promo_file_path(name))
        promo = self._file("paihuo-promo-31-web.mp4")
        self.assertEqual(os.path.realpath(promo), avatar.promo_file_path("paihuo-promo-31-web.mp4"))

    def test_cleanup_removes_stale_files_but_not_active_job_or_library(self):
        library = "1" * 32 + ".jpg"
        active_audio = "2" * 32 + ".mp3"
        active_photo = "3" * 32 + ".png"
        done_tts = "4" * 32 + ".mp3"
        orphan = "5" * 32 + ".jpg"
        fresh = "6" * 32 + ".jpg"
        for name in (library, active_audio, active_photo, done_tts, orphan):
            self._file(name, age_days=40)
        self._file(fresh, age_days=1)
        self._file("paihuo-promo-31-web.mp4", age_days=400)
        self._file(".upload-abc.part", age_days=3)
        db.set_setting("avatar_photos:2", json.dumps([{"name": library}]))
        db.insert("avatar_job", {
            "tenant_id": 2, "status": "running", "billing_status": "charged",
            "params_json": json.dumps({"photo_name": active_photo}),
            "audio_file": f"/files/avatar-public/{active_audio}",
        })
        db.insert("avatar_job", {
            "tenant_id": 2, "status": "done", "billing_status": "succeeded",
            "params_json": json.dumps({"photo_name": library}),
            "audio_file": f"/files/avatar-public/{done_tts}",
        })
        report = avatar.cleanup_public_assets()
        left = set(os.listdir(avatar.PUBLIC_DIR))
        self.assertEqual({done_tts, orphan, ".upload-abc.part"}, set(report["removed"]))
        for name in (library, active_audio, active_photo, fresh, "paihuo-promo-31-web.mp4"):
            self.assertIn(name, left)
        self.assertEqual(2, report["kept_active"])
        self.assertEqual(1, report["kept_library"])
        # 天数可配
        features.set_config({"pub_cleanup_days": 365})
        self._file(orphan, age_days=40)
        self.assertEqual([], avatar.cleanup_public_assets()["removed"])


# ---------------------------------------------------------------- 3. 授权
class ConsentTests(_DbCase):
    def test_missing_consent_rejected_and_recorded_when_given(self):
        photo = "7" * 32 + ".jpg"
        with self.assertRaises(avatar.ConsentRequired) as caught:
            avatar.require_consent(2, [photo], "", "avatar_job", user=auth.current())
        self.assertIn("授权", str(caught.exception))
        self.assertFalse(avatar.has_consent(2, photo))
        before = time.time()
        records = avatar.require_consent(2, [photo, None], True, "avatar_job",
                                         user=auth.current(), kinds={photo: "photo"})
        self.assertEqual(1, len(records))
        entry = avatar.consent_records(2, photo)[0]
        self.assertEqual(20, entry["user_id"])
        self.assertEqual("老板2", entry["username"])
        self.assertEqual(photo, entry["asset"])
        self.assertEqual("photo", entry["kind"])
        self.assertEqual(avatar.CONSENT_VERSION, entry["version"])
        self.assertGreaterEqual(entry["ts"], before)
        self.assertTrue(avatar.has_consent(2, photo))
        # 已有记录后不必重复勾选;其他企业看不到这条记录
        self.assertEqual([], avatar.require_consent(2, [photo], "", "avatar_job"))
        self.assertFalse(avatar.has_consent(3, photo))
        self.assertEqual([], avatar.consent_records(3))

    def test_upload_record_and_old_version_needs_reconsent(self):
        voice = "8" * 32 + ".mp3"
        avatar.record_consent(2, voice, "voice", "upload", user=auth.current())
        self.assertEqual("upload", avatar.consent_records(2)[0]["action"])
        key = f"avatar_consent:2:{voice}"
        items = json.loads(db.get_setting(key))
        items[0]["version"] = "old"
        db.set_setting(key, json.dumps(items))
        self.assertFalse(avatar.has_consent(2, voice))
        with self.assertRaises(ValueError):
            avatar.record_consent(2, "../etc/passwd", "voice", "upload")

    def test_overseas_provider_requires_explicit_opt_in(self):
        with self.assertRaises(avatar.ConsentRequired) as caught:
            avatar.overseas_allowed("heygen", "")
        self.assertIn("境外", str(caught.exception))
        self.assertTrue(avatar.overseas_allowed("heygen", True))
        self.assertFalse(avatar.overseas_allowed("", None))
        self.assertFalse(avatar.overseas_allowed("kling", "0"))


# ---------------------------------------------------------------- 4. AI 标识
class AiLabelTests(_DbCase):
    PACKS = [{"platform": "小红书", "title": "t", "body": "正文一"},
             {"platform": "公众号", "title": "t", "body": "正文二"}]

    def test_label_on_by_default_in_packs_and_idempotent(self):
        packs = features.label_packs(self.PACKS, 2)
        for pack in packs:
            self.assertTrue(pack["body"].endswith(features.DEFAULT_AI_LABEL))
        again = features.label_packs(packs, 2)
        self.assertEqual(1, again[0]["body"].count(features.DEFAULT_AI_LABEL))
        self.assertEqual("正文一", self.PACKS[0]["body"])

    def test_label_text_configurable_per_platform_and_can_be_disabled(self):
        features.save_ai_label_conf(2, {"text": "AI 协助创作", "off_platforms": ["公众号"]})
        packs = features.label_packs(self.PACKS, 2)
        self.assertTrue(packs[0]["body"].endswith("AI 协助创作"))
        self.assertEqual("正文二", packs[1]["body"])
        self.assertEqual("", features.ai_label_for(2, "公众号"))
        # 别的企业不受影响
        self.assertTrue(features.label_packs(self.PACKS, 3)[1]["body"].endswith(
            features.DEFAULT_AI_LABEL))
        features.save_ai_label_conf(2, {"enabled": False})
        self.assertEqual(self.PACKS, features.label_packs(self.PACKS, 2))
        self.assertEqual("# 标题\n\n正文", features.label_markdown("# 标题\n\n正文", 2))
        with self.assertRaises(ValueError):
            features.save_ai_label_conf(2, {"text": "长" * 31})

    def test_exported_docx_contains_label(self):
        from docx import Document
        from app import export

        md = features.label_markdown("# 火锅上新\n\n今天上新毛肚。", 2)
        data = export.md_to_docx(md, "火锅上新")
        text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
        self.assertIn(features.DEFAULT_AI_LABEL, text)
        features.save_ai_label_conf(2, {"enabled": False})
        data = export.md_to_docx(features.label_markdown("# 火锅上新\n\n正文", 2), "火锅上新")
        text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
        self.assertNotIn(features.DEFAULT_AI_LABEL, text)

    def test_wechat_layout_and_video_end_card_labels(self):
        html = mplayout.render("## 小节\n\n正文", "orange", ai_label="本内容由 AI 辅助生成")
        self.assertIn("本内容由 AI 辅助生成", html)
        self.assertNotIn("AI 辅助生成", mplayout.render("## 小节\n\n正文", "orange"))
        end = features.label_end_text(2, "")
        from app import textvideo
        lines = textvideo._wrap(end, 12, 3).split("\n")
        self.assertEqual("本内容由AI辅助生成", lines[-1])
        self.assertTrue(lines[0].startswith("喜欢这条就点赞关注"))
        long_end = features.label_end_text(2, "关注我们每天分享一道家常菜做法简单又好吃还省钱")
        self.assertLessEqual(len(long_end), 40)
        self.assertEqual("本内容由AI辅助生成", textvideo._wrap(long_end, 12, 3).split("\n")[-1])
        self.assertEqual(long_end, features.label_end_text(2, long_end))
        features.save_ai_label_conf(2, {"enabled": False})
        self.assertEqual("关注我们", features.label_end_text(2, "关注我们"))

    @unittest.skipUnless(shutil.which("ffmpeg"), "需要 ffmpeg")
    def test_avatar_video_metadata_gets_label(self):
        path = os.path.join(self.tmp.name, "clip.mp4")
        subprocess.run(
            ["ffmpeg", "-nostdin", "-y", "-loglevel", "error", "-f", "lavfi", "-i",
             "color=c=black:s=64x64:d=0.5", "-pix_fmt", "yuv420p", path],
            check=True, timeout=60)
        self.assertTrue(avatar.stamp_video_label(path, "本内容由 AI 辅助生成"))
        probe = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format_tags=comment,title",
             "-of", "json", path], capture_output=True, text=True, timeout=60)
        tags = json.loads(probe.stdout)["format"]["tags"]
        self.assertEqual("本内容由 AI 辅助生成", tags.get("title"))
        self.assertFalse(avatar.stamp_video_label(os.path.join(self.tmp.name, "none.mp4"), "x"))


def tearDownModule():
    if _DbCase._template_dir:
        shutil.rmtree(_DbCase._template_dir, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
